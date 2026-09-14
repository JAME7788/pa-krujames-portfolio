from pathlib import Path
from copy import deepcopy
import re
import shutil
from io import BytesIO
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
SOURCE=ROOT/'assets/docs/รายงาน PA 2569 ครูอนันตชัย ป1 11คน ฉบับปรับปรุง.docx'
OUT=ROOT/'deliverables/จัดรูปเล่มใหม่-รายงาน PA 2569 ป1 11คน.docx'
FONT='TH SarabunPSK'
doc=Document(SOURCE)
backup=ROOT/'tmp/reformat-pa/original'
backup.mkdir(parents=True,exist_ok=True)
for suffix in ['.docx','.pdf']:
    source=SOURCE.with_suffix(suffix)
    if not (backup/source.name).exists(): shutil.copy2(source,backup/source.name)

def font(run,size=16,bold=None):
    run.font.name=FONT
    run.font.size=Pt(size)
    run.font.color.rgb=RGBColor(0,0,0)
    if bold is not None:run.bold=bold
    pr=run._r.get_or_add_rPr()
    fonts=pr.find(qn('w:rFonts'))
    if fonts is None:fonts=OxmlElement('w:rFonts');pr.append(fonts)
    for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),FONT)

for section in doc.sections:
    section.page_width=Cm(21);section.page_height=Cm(29.7)
    section.top_margin=Cm(2);section.bottom_margin=Cm(1.9)
    section.left_margin=Cm(2.3);section.right_margin=Cm(2)
    section.header_distance=Cm(.8);section.footer_distance=Cm(.8)
for name,size in [('Normal',16),('Title',28),('Subtitle',18),('Heading 1',22),('Heading 2',19),('Heading 3',17),('Caption',12),('TOC 1',16)]:
    if name not in doc.styles:doc.styles.add_style(name,1)
    style=doc.styles[name]
    style.font.name=FONT;style.font.size=Pt(size);style.font.color.rgb=RGBColor(0,0,0)
    style.paragraph_format.line_spacing=1.0
    style.paragraph_format.space_after=Pt(3)
    if name.startswith('Heading'):
        style.font.bold=True;style.paragraph_format.keep_with_next=True
        style.paragraph_format.space_before=Pt(8)

# Remove cached TOC entries, then rebuild the field using the new chapter hierarchy.
paras=list(doc.paragraphs)
contents=next(p for p in paras if p.text=='สารบัญ')
next_chapter=next(p for p in paras[paras.index(contents)+1:] if p.style.name=='Heading 1')
node=contents._p.getnext()
while node is not None and node!=next_chapter._p:
    following=node.getnext();node.getparent().remove(node);node=following
field_p=OxmlElement('w:p')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'TOC \\o "1-1" \\h \\z')
field_p.append(field);contents._p.addnext(field_p)

chapters={
 'ข้อมูลผู้รับการประเมินและขอบเขตรายงาน':'ส่วนที่ 1 ข้อมูลผู้รับการประเมินและภาระงาน',
 'ด้านที่ 1 การจัดการเรียนรู้':'ส่วนที่ 2 ผลการปฏิบัติงานตามมาตรฐานตำแหน่ง',
 'ประเด็นท้าทายและแนวทางพัฒนา':'ส่วนที่ 3 การพัฒนางานตามประเด็นท้าทาย',
 'แผนการจัดการเรียนรู้และกำหนดการ':'ภาคผนวก ก แผนการจัดการเรียนรู้',
 'ใบงานประกอบการฝึกเมาส์ชุดที่ 1':'ภาคผนวก ข ใบงานและเครื่องมือประเมิน',
 'ภาพหลักฐานการนำทักษะเมาส์ไปใช้':'ภาคผนวก ค ภาพและทะเบียนหลักฐาน',
}
appendix=False
challenge_part=False
lesson_intro=False
for p in list(doc.paragraphs):
    text=p.text
    if text=='ประเด็นท้าทายและแนวทางพัฒนา':challenge_part=True
    if text=='แผนการจัดการเรียนรู้และกำหนดการ':appendix=True
    if re.match(r'^แผนที่ [5-8] ',text):lesson_intro=True
    if text.startswith(('สื่อและการประเมินแผนที่','ใบงาน')):lesson_intro=False
    if text in chapters:
        p.text=chapters[text];p.style='Heading 1';p.paragraph_format.page_break_before=True
    elif p.style.name=='Heading 1' and text!='คำนำและสรุปผลการดำเนินงาน':
        p.style='Heading 2'
        p.paragraph_format.page_break_before=bool(text.startswith(('ภาระงานและ','แผนที่ ','ใบงาน','แบบทดสอบ','เกณฑ์ประเมิน','บัญชี','แบบสังเกต','แบบสะท้อน')))
    if re.match(r'^[123]\.\d+ ',text):p.style='Heading 3'
    if text in ['ผลการพัฒนาและการต่อยอด','การควบคุมตัวชี้และท่าจับเมาส์','แนบช่วยเหลือหรือภารกิจต่อยอด','แผนช่วยเหลือหรือภารกิจต่อยอด','ทะเบียนหลักฐานและแหล่งข้อมูล']:
        p.paragraph_format.page_break_before=True
    if p.style.name in ['Normal','Caption']:
        p.paragraph_format.widow_control=True
        p.paragraph_format.space_after=Pt(3)
        p.paragraph_format.line_spacing=1.0
    for run in p.runs:
        size=12 if p.style.name=='Caption' else (16 if lesson_intro else (15 if appendix or challenge_part else 16))
        if p.style.name.startswith('Heading'):size={'Heading 1':22,'Heading 2':19,'Heading 3':17}[p.style.name]
        if p.style.name in ['Title','Subtitle']:size=28 if p.style.name=='Title' else 18
        font(run,size)

def format_table(table,labelled=False):
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    pr=table._tbl.tblPr
    for old in list(pr.findall(qn('w:tblBorders'))):pr.remove(old)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        elem=OxmlElement('w:'+edge);elem.set(qn('w:val'),'single');elem.set(qn('w:sz'),'4');elem.set(qn('w:color'),'CBDDEB');borders.append(elem)
    pr.append(borders)
    for ri,row in enumerate(table.rows):
        rowpr=row._tr.get_or_add_trPr()
        for node in list(rowpr.findall(qn('w:trHeight'))):rowpr.remove(node)
        if rowpr.find(qn('w:cantSplit')) is None:rowpr.append(OxmlElement('w:cantSplit'))
        if ri==0 and not labelled and rowpr.find(qn('w:tblHeader')) is None:rowpr.append(OxmlElement('w:tblHeader'))
        for ci,cell in enumerate(row.cells):
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
            tcpr=cell._tc.get_or_add_tcPr()
            for old in list(tcpr.findall(qn('w:shd')))+list(tcpr.findall(qn('w:tcMar'))):tcpr.remove(old)
            shd=OxmlElement('w:shd')
            fill=('EAF3FA' if ci==0 else 'FFFFFF') if labelled else ('D6E9F7' if ri==0 else ('F4F8FC' if ri%2==0 else 'FFFFFF'))
            shd.set(qn('w:fill'),fill);tcpr.append(shd)
            margins=OxmlElement('w:tcMar')
            vertical='35' if len(table.rows)>=10 else '60'
            for key,value in [('top',vertical),('bottom',vertical),('left','100'),('right','100')]:
                el=OxmlElement('w:'+key);el.set(qn('w:w'),value);el.set(qn('w:type'),'dxa');margins.append(el)
            tcpr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing=1.0;p.paragraph_format.space_after=Pt(2)
                p.paragraph_format.widow_control=True
                if labelled:p.paragraph_format.keep_with_next=ri<len(table.rows)-1
                for run in p.runs:font(run,13 if len(table.rows)>=10 else 14,bold=(ci==0 if labelled else ri==0))

# Present every indicator using the reference's task / result / evidence structure.
indicators=0
for heading in list(doc.paragraphs):
    if not re.match(r'^[123]\.\d+ ',heading.text):continue
    body=[];node=heading._p.getnext()
    while node is not None and node.tag==qn('w:p'):
        style=node.find('w:pPr/w:pStyle',node.nsmap)
        if style is not None and ('Heading' in style.get(qn('w:val'),'') or style.get(qn('w:val'))=='ContentsHeading'):break
        body.append(node);node=node.getnext()
    from docx.text.paragraph import Paragraph
    blocks=[Paragraph(n,doc).text for n in body if Paragraph(n,doc).text.strip()]
    if not blocks:continue
    tasks=[x for x in blocks if not x.startswith(('ผลที่เกิดขึ้น','หลักฐานประกอบ'))]
    outcomes=[x for x in blocks if x.startswith('ผลที่เกิดขึ้น')]
    refs=[x for x in blocks if x.startswith('หลักฐานประกอบ')]
    table=doc.add_table(rows=3,cols=2)
    for col,width in zip(table.columns,[3.2,13.5]):col.width=Cm(width)
    for row,label,values in zip(table.rows,['การดำเนินงาน','ผลและการติดตาม','หลักฐานอ้างอิง'],[tasks,outcomes,refs]):
        row.cells[0].width=Cm(3.2);row.cells[1].width=Cm(13.5)
        row.cells[0].text=label;row.cells[1].text='\n'.join(values)
    heading._p.addnext(table._tbl)
    for n in body:n.getparent().remove(n)
    format_table(table,True);indicators+=1
assert indicators==15,indicators
for table in doc.tables:
    if len(table.rows)==3 and table.cell(0,0).text=='การดำเนินงาน':continue
    format_table(table)

# Balance the cover without introducing unsupported seals or award graphics.
first=doc.paragraphs[0]
first.paragraph_format.space_before=Pt(18)
first.paragraph_format.space_after=Pt(14)
cover_paras=list(doc.paragraphs)
name=next(p for p in cover_paras if p.text=='นายอนันตชัย เพ็ชรรี่')
portrait=doc.add_paragraph()
portrait.alignment=WD_ALIGN_PARAGRAPH.CENTER
portrait_image=BytesIO()
Image.open(ROOT/'assets/images/profile-suit-krujames.jpg').save(portrait_image,format='PNG')
portrait_image.seek(0)
portrait.add_run().add_picture(portrait_image,width=Cm(3.3))
portrait.paragraph_format.space_before=Pt(8);portrait.paragraph_format.space_after=Pt(12)
name._p.addprevious(portrait._p)
for p in list(doc.paragraphs):
    if p.text=='คำนำและสรุปผลการดำเนินงาน':break
    p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next=True
    if p.text=='ประเด็นท้าทาย':p.paragraph_format.space_before=Pt(12)
for section in doc.sections:
    section.different_first_page_header_footer=True
    for p in section.header.paragraphs:
        for run in p.runs:font(run,11)
    for p in section.footer.paragraphs:
        for run in p.runs:font(run,12)
OUT.parent.mkdir(exist_ok=True)
doc.save(OUT)
print(f'{OUT}\nIndicators: {indicators}; tables: {len(doc.tables)}')
