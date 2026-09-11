from pathlib import Path
from PIL import Image, ImageDraw
from pypdf import PdfReader
from docx import Document
import json
import re

root=Path(__file__).resolve().parents[1]
render=root/'qa-pa/updated-render'
pdf=PdfReader(render/'report.pdf')
document=Document(root/'deliverables/รายงาน PA 2569 ครูอนันตชัย ป1 11คน ฉบับปรับปรุง.docx')
text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(c.text for t in document.tables for row in t.rows for c in row.cells)
checks={
    'pages':len(pdf.pages),
    'sparse_pages':[(i+1,len(p.extract_text())) for i,p in enumerate(pdf.pages) if len(p.extract_text())<350],
    'forbidden_template_residue':{s:len(re.findall(r'(?<!\d)'+re.escape(s),text)) for s in ['คณิตศาสตร์','จินตนา','ดวงกมล','42.22','88.89','17.78','8.44','9 คน','๙ คน','assets_drop','t-test']},
    'tables':len(document.tables),
    'photos':len(document.inline_shapes),
}
(render/'audit.json').write_text(json.dumps(checks,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(checks,ensure_ascii=False,indent=2))
for i in range(0,len(pdf.pages),2):
    files=[render/f'final-{n:02d}.png' for n in range(i+1,min(i+3,len(pdf.pages)+1))]
    images=[Image.open(file).convert('RGB') for file in files]
    w=sum(im.width for im in images)+30*(len(images)+1)
    h=max(im.height for im in images)+65
    sheet=Image.new('RGB',(w,h),'#DDE1E5')
    x=30
    for n,im in enumerate(images,i+1):
        sheet.paste(im,(x,35))
        ImageDraw.Draw(sheet).text((x,10),f'PAGE {n}',fill='black')
        x+=im.width+30
    sheet.save(render/f'review-{i//2+1:02d}.png')
