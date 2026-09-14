from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
source = ROOT/'deliverables/จัดรูปเล่มใหม่-รายงาน PA 2569 ป1 11คน.docx'
out = ROOT/'deliverables/รายงาน PA 2569 ฉบับตกแต่ง.docx'
doc = Document(source)

def shade(pr, color):
    for old in list(pr.findall(qn('w:shd'))):
        pr.remove(old)
    node = OxmlElement('w:shd')
    node.set(qn('w:fill'), color)
    pr.append(node)

def rule(p, edge, color, size):
    pr = p._p.get_or_add_pPr()
    borders = pr.find(qn('w:pBdr'))
    if borders is None:
        borders = OxmlElement('w:pBdr')
        pr.append(borders)
    node = OxmlElement('w:'+edge)
    for key, value in [('val','single'),('sz',str(size)),('color',color),('space','5')]:
        node.set(qn('w:'+key), value)
    borders.append(node)

for p in doc.paragraphs:
    if p.style.name == 'Heading 1':
        shade(p._p.get_or_add_pPr(), 'E7F2FC')
    elif p.style.name == 'Heading 3':
        shade(p._p.get_or_add_pPr(), 'F1F7FC')

for table in doc.tables:
    labelled = table.cell(0,0).text == 'การดำเนินงาน'
    for ri,row in enumerate(table.rows):
        for ci,cell in enumerate(row.cells):
            header = ci == 0 if labelled else ri == 0
            shade(cell._tc.get_or_add_tcPr(), ('D9ECFA' if labelled else '17548C') if header else ('F1F7FC' if ri%2 == 0 else 'FFFFFF'))
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.color.rgb = RGBColor.from_string('FFFFFF' if header and not labelled else '000000')

for section in doc.sections:
    for p in section.header.paragraphs:
        rule(p, 'bottom', '2674AE', 10)
    for p in section.footer.paragraphs:
        rule(p, 'top', 'C4DCEE', 4)
    p = section.first_page_header.paragraphs[0]
    p.text = 'PA 2569    |    รายงานผลการพัฒนางานตามข้อตกลง'
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shade(p._p.get_or_add_pPr(), 'D9ECFA')
    for run in p.runs:
        run.font.name = 'TH SarabunPSK'
        run.font.size = Pt(16)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0,0,0)
    p = section.first_page_footer.paragraphs[0]
    p.text = 'โรงเรียนบ้านคลองมดแดง'
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in p.runs:
        run.font.name = 'TH SarabunPSK'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0,0,0)
    rule(p, 'top', '2674AE', 16)

doc.save(out)
print(out)
