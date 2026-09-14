from pathlib import Path
import json
import re
from io import BytesIO

from PIL import Image, ImageOps
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "deliverables/รายงาน PA 2569 ฉบับตกแต่ง.docx"
OUTPUT = ROOT / "deliverables/รายงาน PA 2569 พร้อมภาพรายตัวชี้วัด.docx"
MANIFEST = ROOT / "tmp/indicator-evidence/manifest.json"

# One primary item per indicator. Paths are intentionally unique so the report
# does not recycle the same photograph under several criteria.
EVIDENCE = {
    "1.1": ("assets/evidence_pa_photos/IMG_20260910_133433.jpg", "การนำหน่วยการเรียนรู้วิทยาการคำนวณตามหลักสูตรสถานศึกษาไปใช้จริง", "10 กันยายน 2569"),
    "1.2": ("assets/evidence_pa_photos/MVIMG_20260514_151034.jpg", "การจัดกิจกรรมตามแผน Active Learning โดยครูสาธิตและให้ผู้เรียนลงมือปฏิบัติ", "14 พฤษภาคม 2569"),
    "1.3": ("assets/evidence_pa_photos/IMG_20260716_114228.jpg", "ครูให้คำแนะนำรายกลุ่มระหว่างกิจกรรม Coding และการสร้างชิ้นงานดิจิทัล", "16 กรกฎาคม 2569"),
    "1.4": ("assets/evidence_pa_photos/IMG_20260803_084731.jpg", "การใช้สื่อเกมดิจิทัลและ Gamification เพื่อกระตุ้นการเรียนรู้", "3 สิงหาคม 2569"),
    "1.5": ("assets/evidence_pa_photos/IMG_20260611_132439.jpg", "ผู้เรียนทำแบบประเมินและภารกิจดิจิทัลรายบุคคลในห้องคอมพิวเตอร์", "11 มิถุนายน 2569"),
    "1.6": ("assets/evidence_pa_photos/IMG_20260128_140738.jpg", "การสังเกตปัญหา ให้คำแนะนำเฉพาะจุด และช่วยเหลือผู้เรียนระหว่างปฏิบัติ", "28 มกราคม 2569"),
    "1.7": ("assets/evidence_pa_photos/IMG_20260910_111013.jpg", "บรรยากาศห้องปฏิบัติการที่เอื้อต่อการเรียนรู้แบบหนึ่งคนหนึ่งเครื่อง", "10 กันยายน 2569"),
    "1.8": ("assets/evidence_pa_photos/IMG_20260911_075634.jpg", "กิจกรรมส่งเสริมสมาธิ คุณธรรม และคุณลักษณะอันพึงประสงค์ของผู้เรียน", "11 กันยายน 2569"),
    "2.1": ("assets/evidence_pa_photos/MVIMG_20260518_132255.jpg", "การใช้ระบบดิจิทัลบันทึกข้อมูลและติดตามความก้าวหน้าของผู้เรียนรายบุคคล", "18 พฤษภาคม 2569"),
    "2.2": ("assets/evidence_pa_photos/MVIMG_20260519_164336.jpg", "การเยี่ยมบ้านเพื่อศึกษาสภาพจริงและวางแนวทางช่วยเหลือผู้เรียน", "19 พฤษภาคม 2569"),
    "2.3": ("assets/evidence_pa_photos/IMG_20260224_091508.jpg", "การดูแลระบบคอมพิวเตอร์และโครงสร้างพื้นฐานเพื่อสนับสนุนงานของสถานศึกษา", "24 กุมภาพันธ์ 2569"),
    "2.4": ("assets/evidence_pa_photos/MVIMG_20260506_085132.jpg", "การประชุมผู้ปกครองเพื่อสื่อสารข้อมูลผู้เรียนและสร้างความร่วมมือบ้านกับโรงเรียน", "6 พฤษภาคม 2569"),
    "3.1": ("assets_drop/06_ภาพเกียรติบัตรและรางวัล/LINE_ALBUM_อนันตชัย_260712_1.jpg", "เกียรติบัตรการพัฒนาความรู้ด้าน AI Learning Hub และ OBEC Content Center", "7 มิถุนายน 2569"),
    "3.2": ("assets/evidence_pa_photos/MVIMG_20260506_155728.jpg", "การประชุมแลกเปลี่ยนเรียนรู้ทางวิชาชีพและร่วมวางแผนพัฒนาผู้เรียน", "6 พฤษภาคม 2569"),
    "3.3": ("assets/evidence_pa_photos/IMG_20260724_085632.jpg", "การนำความรู้ด้านสื่อสามมิติและเกมดิจิทัลมาประยุกต์ใช้ในการจัดการเรียนรู้", "24 กรกฎาคม 2569"),
}


def prepared_image(path: Path) -> tuple[BytesIO, float]:
    with Image.open(path) as source:
        image = ImageOps.exif_transpose(source).convert("RGB")
        image.thumbnail((2200, 1650), Image.Resampling.LANCZOS)
        stream = BytesIO()
        image.save(stream, format="JPEG", quality=91, optimize=True)
        stream.seek(0)
        return stream, min(15.7, 11.2 * image.width / image.height)


doc = Document(SOURCE)
manifest = []

for paragraph in doc.paragraphs:
    if paragraph.text == "ทะเบียนการพัฒนาตนเองและงานชุมชน":
        paragraph.paragraph_format.page_break_before = True

for heading in list(doc.paragraphs):
    match = re.match(r"^([123]\.\d+) ", heading.text)
    if not match:
        continue
    key = match.group(1)
    table = heading._p.getnext()
    if table is None or table.tag != qn("w:tbl"):
        raise RuntimeError(f"ไม่พบตารางรายละเอียดของตัวชี้วัด {key}")

    previous_section = None
    sibling = heading._p.getprevious()
    while sibling is not None:
        if sibling.tag == qn("w:p"):
            candidate = Paragraph(sibling, doc)
            if candidate.text.strip():
                previous_section = candidate
                break
        sibling = sibling.getprevious()
    if previous_section is not None and previous_section.text.startswith("ด้านที่ "):
        previous_section.paragraph_format.page_break_before = True
        previous_section.paragraph_format.keep_with_next = True
        heading.paragraph_format.page_break_before = False
    else:
        heading.paragraph_format.page_break_before = True
    relative_path, description, date_text = EVIDENCE[key]
    path = ROOT / relative_path
    if not path.exists():
        raise FileNotFoundError(path)

    anchor = table
    caption = doc.add_paragraph(f"ภาพหลักฐานประกอบตัวชี้วัด {key}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.keep_with_next = True
    for run in caption.runs:
        run.font.name = "TH SarabunPSK"
        run.font.size = Pt(15)
        run.bold = True
        run.font.color.rgb = RGBColor(14, 52, 96)
    anchor.addnext(caption._p)
    anchor = caption._p

    stream, width_cm = prepared_image(path)
    picture = doc.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.paragraph_format.keep_with_next = True
    picture.paragraph_format.space_after = Pt(2)
    picture.add_run().add_picture(stream, width=Cm(width_cm))
    anchor.addnext(picture._p)
    anchor = picture._p

    detail = doc.add_paragraph(f"{description} | วันที่ {date_text}")
    detail.alignment = WD_ALIGN_PARAGRAPH.CENTER
    detail.paragraph_format.space_after = Pt(6)
    for run in detail.runs:
        run.font.name = "TH SarabunPSK"
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(0, 0, 0)
    anchor.addnext(detail._p)

    manifest.append({
        "indicator": key,
        "source": relative_path,
        "caption": description,
        "date": date_text,
    })

if len(manifest) != 15:
    raise RuntimeError(f"แนบภาพได้ {len(manifest)} ตัวชี้วัด แทนที่จะเป็น 15")
if len({item["source"] for item in manifest}) != len(manifest):
    raise RuntimeError("พบภาพซ้ำข้ามตัวชี้วัด")

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
MANIFEST.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
MANIFEST.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
print(f"Added unique evidence for {len(manifest)} indicators: {OUTPUT}")
