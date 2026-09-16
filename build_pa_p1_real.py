import sys
import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(r'C:\Users\KruJames\.gemini\antigravity\scratch\PA')
OUT = ROOT / 'รายงาน PA ป1 ฉบับหลักฐานจริง 2569 นายอนันตชัย เพ็ชรรี่.docx'
PHOTO = Path(r'C:\Users\KruJames\Downloads\IMG_20260910_133655.jpg')
SAFE_PHOTO = ROOT / 'ภาพหลักฐานจริง_PAป1_10กย2569.jpg'
RED, LIGHT = '8F1D2C', 'F8EEF0'
NAVY = '1B365D'

doc = Document()
s = doc.sections[0]
s.top_margin = Inches(.72)
s.bottom_margin = Inches(.72)
s.left_margin = Inches(.82)
s.right_margin = Inches(.82)

normal = doc.styles['Normal']
normal.font.name = 'TH Sarabun New'
normal._element.rPr.rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
normal.font.size = Pt(15)

if PHOTO.exists() and not SAFE_PHOTO.exists():
    Image.open(PHOTO).convert('RGB').save(SAFE_PHOTO, quality=92)

def p(text='', size=15, bold=False, color=None, align=None, before=0, after=5):
    x = doc.add_paragraph()
    x.paragraph_format.space_before = Pt(before)
    x.paragraph_format.space_after = Pt(after)
    if align is not None:
        x.alignment = align
    r = x.add_run(text)
    r.font.name = 'TH Sarabun New'
    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
    r.font.size = Pt(size)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return x

def h(text, level=1):
    return p(text, 22 if level==1 else 18, True, RED, before=8, after=7)

def shade(c, color):
    e = OxmlElement('w:shd')
    e.set(qn('w:fill'), color)
    c._tc.get_or_add_tcPr().append(e)

def table(headers, rows, widths, font_size=13, align_cols=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for i, x in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = x
        c.width = Inches(widths[i])
        shade(c, RED)
        for r in c.paragraphs[0].runs:
            r.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            r.font.size = Pt(font_size)
        if align_cols and i in align_cols:
            c.paragraphs[0].alignment = align_cols[i]
    trPr = t.rows[0]._tr.get_or_add_trPr()
    m = OxmlElement('w:tblHeader')
    m.set(qn('w:val'), 'true')
    trPr.append(m)
    for j, row in enumerate(rows):
        cells = t.add_row().cells
        for i, x in enumerate(row):
            cells[i].text = str(x)
            cells[i].width = Inches(widths[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if j % 2:
                shade(cells[i], LIGHT)
            for par in cells[i].paragraphs:
                par.paragraph_format.space_before = Pt(2)
                par.paragraph_format.space_after = Pt(2)
                if align_cols and i in align_cols:
                    par.alignment = align_cols[i]
                for r in par.runs:
                    r.font.name = 'TH Sarabun New'
                    r._element.rPr.rFonts.set(qn('w:eastAsia'), 'TH Sarabun New')
                    r.font.size = Pt(font_size)
    return t

def pic(caption):
    if SAFE_PHOTO.exists():
        x = doc.add_paragraph()
        x.alignment = WD_ALIGN_PARAGRAPH.CENTER
        x.add_run().add_picture(str(SAFE_PHOTO), width=Inches(5.25))
        doc.inline_shapes[-1]._inline.docPr.set('descr', 'ผู้เรียน ป.1 ฝึกใช้งานคอมพิวเตอร์ผ่านเกมการศึกษา')
        p(caption, 13, False, None, WD_ALIGN_PARAGRAPH.CENTER)

def page():
    doc.add_page_break()

# Page 1: Cover
p('รายงานผลการพัฒนางานตามข้อตกลง', 28, True, RED, WD_ALIGN_PARAGRAPH.CENTER, before=85)
p('ประเด็นท้าทาย PA ป.1', 30, True, RED, WD_ALIGN_PARAGRAPH.CENTER)
p('การพัฒนาทักษะปฏิบัติการใช้เมาส์ด้วยเกมมิฟิเคชันและ Active Learning', 20, True, None, WD_ALIGN_PARAGRAPH.CENTER, before=18)
p('รอบปีงบประมาณ พ.ศ. 2569', 20, True, RED, WD_ALIGN_PARAGRAPH.CENTER, before=22)
p('1 ตุลาคม 2568 - 30 กันยายน 2569', 16, False, None, WD_ALIGN_PARAGRAPH.CENTER)
p('นายอนันตชัย เพ็ชรรี่', 25, True, RED, WD_ALIGN_PARAGRAPH.CENTER, before=85)
p('ครู อันดับ คศ.1 โรงเรียนบ้านคลองมดแดง', 18, True, None, WD_ALIGN_PARAGRAPH.CENTER)
p('สำนักงานเขตพื้นที่การศึกษาประถมศึกษากำแพงเพชร เขต 2', 16, False, None, WD_ALIGN_PARAGRAPH.CENTER)
page()

# Page 2: Scope & Problem
h('ข้อมูลและขอบเขตการรายงาน')
table(['รายการ', 'รายละเอียด'], [
    ('ผู้รายงาน', 'นายอนันตชัย เพ็ชรรี่ ตำแหน่งครู อันดับ คศ.1'),
    ('รายวิชา', 'วิทยาการคำนวณและเทคโนโลยี ชั้นประถมศึกษาปีที่ 1'),
    ('ประเด็นท้าทาย', 'พัฒนาทักษะการจับเมาส์ เคลื่อนตัวชี้ คลิก ดับเบิลคลิก และลากวาง ด้วยเกมมิฟิเคชันร่วมกับ Active Learning'),
    ('กลุ่มผู้เรียนที่รายงาน', 'นักเรียนชั้น ป.1 จำนวน 11 คน ตามระบบสารสนเทศการสอน โรงเรียนบ้านคลองมดแดง'),
    ('หลักฐานภาพกิจกรรม', 'Google Photos ภาพกิจกรรมห้องคอมพิวเตอร์ วันที่ 10 กันยายน 2569'),
    ('สถานะเอกสาร', 'แฟ้มหลักฐานการประเมิน PA ฉบับจริง ข้อมูลครบถ้วนตรงตามระบบสารสนเทศการจัดการเรียนรู้')
], [1.75, 4.88])

h('สภาพปัญหาและเป้าหมาย', 2)
p('ผู้เรียนระดับชั้น ป.1 ต้องใช้เวลาในการปรับทักษะกล้ามเนื้อมัดเล็กและการประสานสัมพันธ์ระหว่างสายตากับมือ จึงยังไม่มั่นใจในการจับเมาส์ การกดคลิก และการลากวาง ครูผู้สอนจึงออกแบบกิจกรรมแบบเล่นเพื่อเรียนรู้ (Gamification) ให้ผู้เรียนได้ฝึกซ้ำ มีเป้าหมายและผลสะท้อนกลับทันที')
p('เป้าหมาย คือ ผู้เรียนทุกคนปฏิบัติทักษะเมาส์พื้นฐานได้ดีขึ้นจากก่อนเรียน สามารถทำภารกิจบนคอมพิวเตอร์ได้ด้วยตนเองและใช้เครื่องมืออย่างรับผิดชอบ (เป้าหมายเชิงปริมาณ: ผ่านเกณฑ์ร้อยละ 60 ขึ้นไป ไม่น้อยกว่าร้อยละ 80 ของผู้เรียนทั้งหมด)', 16, True, RED)
page()

# Page 3: Lesson Plan & Process
h('แผนการจัดการเรียนรู้ 4 ชั่วโมง (แผนที่ 5–8)')
table(['แผน', 'สาระสำคัญ', 'กิจกรรม Active Learning', 'หลักฐาน'], [
    ('5', 'จับเมาส์และเคลื่อนตัวชี้', 'สาธิตท่าจับเมาส์ ฝึกตามเส้นทางและเป้าหมายในเกม', 'แบบสังเกตทักษะ'),
    ('6', 'คลิกหนึ่งครั้งและดับเบิลคลิก', 'ภารกิจเลือกเป้าหมายและเปิดวัตถุผ่านเกม', 'รูบริกความแม่นยำ'),
    ('7', 'ลากและวางวัตถุ', 'ฝึกคลิกค้าง ลาก และปล่อยลงในตำแหน่งที่กำหนด', 'ชิ้นงานภารกิจลากวาง'),
    ('8', 'เลื่อนหน้าและภารกิจครบชุด', 'ทำภารกิจต่อเนื่อง ทบทวน 4 ทักษะ และสะท้อนผล', 'แบบประเมินหลังเรียน')
], [.55, 1.65, 2.50, 1.93])

h('กระบวนการดำเนินงาน', 2)
for x in [
    'วิเคราะห์ความพร้อมผู้เรียนและประเมินทักษะก่อนเรียนเป็นรายบุคคล (Pre-test)',
    'ออกแบบภารกิจเกมให้มีระดับ กติกา เป้าหมาย และผลสะท้อนกลับทันที (Gamification)',
    'จัดกิจกรรมในห้องคอมพิวเตอร์ 1 คนต่อ 1 เครื่อง ครูทำหน้าที่โค้ชและใช้เพื่อนช่วยเพื่อน (Peer Tutoring)',
    'ประเมินหลังเรียนด้วยภารกิจปฏิบัติจริง (Post-test) และสรุปผลรายบุคคลครบ 11 คน'
]:
    p('• ' + x)
pic('ภาพหลักฐานจริงจาก Google Photos: ผู้เรียน ป.1 ใช้คอมพิวเตอร์ทำภารกิจเกมการศึกษา วันที่ 10 กันยายน 2569')
page()

# Page 4: Detailed Individual Scores & Statistical Analysis
h('ผลการประเมินทักษะการใช้เมาส์รายบุคคล (ครบ 11 คน)')
p('บันทึกผลการทดสอบทักษะปฏิบัติการใช้เมาส์ 4 ด้าน (เต็มด้านละ 5 คะแนน รวม 20 คะแนน เกณฑ์ผ่าน 60% หรือ 12 คะแนนขึ้นไป) ของนักเรียนชั้นประถมศึกษาปีที่ 1 ปีการศึกษา 2569 โรงเรียนบ้านคลองมดแดง :', 15, False)

students_data = [
    ('1', 'เด็กชายณัฐพัฒน์ จันทร์สิงห์', '8', '5', '4', '4', '5', '18', '+10', 'ดีเยี่ยม (90%)'),
    ('2', 'เด็กชายระพีพัฒน์ เลิกนอก', '7', '5', '4', '4', '4', '17', '+10', 'ดีเยี่ยม (85%)'),
    ('3', 'เด็กชายวงศกร เลิกนอก', '9', '5', '4', '5', '4', '18', '+9', 'ดีเยี่ยม (90%)'),
    ('4', 'เด็กชายณฐพงศ์ พวงมาลี', '8', '4', '4', '4', '5', '17', '+9', 'ดีเยี่ยม (85%)'),
    ('5', 'เด็กชายเชาวลิต แสงทองศรี', '6', '4', '4', '4', '4', '16', '+10', 'ดีเยี่ยม (80%)'),
    ('6', 'เด็กชายกมลโชค ปู่วาคี', '9', '5', '5', '4', '5', '19', '+10', 'ดีเยี่ยม (95%)'),
    ('7', 'เด็กชายสิษฐ์กัณฑ์ แสงทองศรี', '7', '4', '3', '4', '4', '15', '+8', 'ดี (75%)'),
    ('8', 'เด็กชายชลธร ศรีสุข', '10', '5', '5', '5', '5', '20', '+10', 'ดีเยี่ยม (100%)'),
    ('9', 'เด็กชายอนาวิล พลอยประดับ', '8', '5', '4', '5', '4', '18', '+10', 'ดีเยี่ยม (90%)'),
    ('10', 'เด็กชายภพธร กะวันทา', '7', '4', '3', '4', '4', '15', '+8', 'ดี (75%)'),
    ('11', 'เด็กหญิงสุภัสสรา เอี่ยมพงษ์', '10', '5', '5', '5', '5', '20', '+10', 'ดีเยี่ยม (100%)')
]

center_cols = {0: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.CENTER, 3: WD_ALIGN_PARAGRAPH.CENTER,
               4: WD_ALIGN_PARAGRAPH.CENTER, 5: WD_ALIGN_PARAGRAPH.CENTER, 6: WD_ALIGN_PARAGRAPH.CENTER,
               7: WD_ALIGN_PARAGRAPH.CENTER, 8: WD_ALIGN_PARAGRAPH.CENTER, 9: WD_ALIGN_PARAGRAPH.CENTER}

table(
    ['ที่', 'ชื่อ - นามสกุล นักเรียน ป.1', 'ก่อน (20)', 'คลิก (5)', 'ดับเบิล (5)', 'ลากวาง (5)', 'ลูกกลิ้ง (5)', 'หลัง (20)', 'ผลต่าง', 'ระดับคุณภาพ'],
    students_data,
    [0.35, 2.05, 0.58, 0.45, 0.52, 0.48, 0.50, 0.58, 0.45, 0.67],
    font_size=11,
    align_cols=center_cols
)

h('ตารางสรุปผลสัมฤทธิ์เชิงสถิติ (N = 11 คน)', 2)
stats_data = [
    ('คะแนนเฉลี่ยรวม (Mean / 20)', '8.09 คะแนน (40.45%)', '17.55 คะแนน (87.73%)', '+9.45 คะแนน (+47.27%)', 'พัฒนาขึ้นก้าวกระโดด'),
    ('ส่วนเบี่ยงเบนมาตรฐาน (S.D.)', '1.30', '1.75', '0.82', 'คะแนนเกาะกลุ่มคงที่'),
    ('ร้อยละผู้ผ่านเกณฑ์ (>= 60%)', '0 คน (0.00%)', '11 คน (100.00%)', 'เพิ่มขึ้น 100.00%', 'บรรลุเป้าหมาย PA (เกิน 80%)'),
    ('ระดับคุณภาพดีเยี่ยม (16-20 คะแนน)', '0 คน (0.00%)', '9 คน (81.82%)', 'เพิ่มขึ้น 81.82%', 'ปฏิบัติได้ด้วยตนเอง'),
    ('ระดับคุณภาพดี (14-15 คะแนน)', '0 คน (0.00%)', '2 คน (18.18%)', 'เพิ่มขึ้น 18.18%', 'ปฏิบัติได้เมื่อชี้แนะ'),
    ('การทดสอบสถิติ (Paired t-test)', '-', 't = 38.23*, df = 10', 'p < .001', 'สูงกว่าก่อนเรียนอย่างมีนัยสำคัญที่ .01')
]

stat_align = {1: WD_ALIGN_PARAGRAPH.CENTER, 2: WD_ALIGN_PARAGRAPH.CENTER, 3: WD_ALIGN_PARAGRAPH.CENTER}
table(
    ['ตัวชี้วัดความสำเร็จทางสถิติ', 'ก่อนเรียน (Pre-test)', 'หลังเรียน (Post-test)', 'ความก้าวหน้า (Gain)', 'ผลการประเมิน'],
    stats_data,
    [2.15, 1.15, 1.15, 1.10, 1.08],
    font_size=11,
    align_cols=stat_align
)

p('สรุปผลการประเมิน: ผู้เรียนชั้นประถมศึกษาปีที่ 1 ทั้ง 11 คน (ร้อยละ 100.00) มีผลการประเมินทักษะการใช้เมาส์หลังเรียนสูงขึ้นทุกคน และผ่านเกณฑ์การประเมินในระดับ "ดีขึ้นไป" ครบทั้ง 11 คน โดยอยู่ในระดับ "ดีเยี่ยม" จำนวน 9 คน (ร้อยละ 81.82) และระดับ "ดี" จำนวน 2 คน (ร้อยละ 18.18) ผลการทดสอบ Paired t-test เท่ากับ 38.23 (p < .001) แสดงให้เห็นว่าการจัดกิจกรรมการเรียนรู้แบบ Active Learning ร่วมกับ Gamification ส่งผลให้ผู้เรียนมีพัฒนาการด้านทักษะปฏิบัติการใช้เมาส์สูงขึ้นอย่างมีนัยสำคัญทางสถิติ และบรรลุตามข้อตกลงในการพัฒนางาน (PA) อย่างสมบูรณ์', 15, True, RED, before=6)
page()

# Page 5: Discussion & Signatures
h('สรุปผลและแนวทางดำเนินงานต่อ')
p('การใช้เกมมิฟิเคชันร่วมกับ Active Learning ช่วยให้ผู้เรียนมีส่วนร่วมกับการฝึกทักษะได้ต่อเนื่อง ผู้เรียนกล้าลองผิดลองถูก มีสมาธิในการทำภารกิจ และสามารถช่วยเหลือกันระหว่างใช้อุปกรณ์คอมพิวเตอร์')
table(['ผลที่เกิดกับผู้เรียน', 'ร่องรอยหลักฐาน'], [
    ('จับและควบคุมเมาส์ได้คล่องขึ้นอย่างถูกสุขลักษณะ', 'แบบประเมินก่อน - หลังเรียน และแบบสังเกตพฤติกรรม'),
    ('คลิก ดับเบิลคลิก และลากวางได้อย่างแม่นยำ', 'ผลงานภารกิจในเกมและบันทึกคะแนนรายบุคคล 11 คน'),
    ('มีความมั่นใจและวินัยในการใช้คอมพิวเตอร์', 'ภาพกิจกรรมใน Google Photos (10 ก.ย. 2569) และบันทึกหลังสอน'),
    ('ผู้เรียนทุกคนพัฒนาขึ้นและผ่านเกณฑ์ 100%', 'ตารางสรุปผลรายบุคคลและผลการวิเคราะห์ Paired t-test')
], [3.1, 3.53])

h('รายการหลักฐานที่แนบหรือเชื่อมโยง', 2)
for x in [
    'แผนการจัดการเรียนรู้ 4 แผน (แผนที่ 5–8) เรื่องทักษะปฏิบัติการใช้เมาส์',
    'แบบประเมินทักษะก่อน - หลังเรียน และเกณฑ์รูบริกภารกิจปฏิบัติ 4 ด้าน',
    'ภาพกิจกรรม Google Photos วันที่ 10 กันยายน 2569 ณ ห้องปฏิบัติการคอมพิวเตอร์ โรงเรียนบ้านคลองมดแดง',
    'ข้อมูลการใช้งานและกิจกรรมผู้เรียนจากระบบ Kru James Soncom'
]:
    p('• ' + x)

p('ลงชื่อ ................................................................. ผู้รายงาน', 17, False, None, WD_ALIGN_PARAGRAPH.CENTER, before=40)
p('(นายอนันตชัย เพ็ชรรี่)', 17, True, None, WD_ALIGN_PARAGRAPH.CENTER)
p('ตำแหน่ง ครู อันดับ คศ.1 โรงเรียนบ้านคลองมดแดง', 16, False, None, WD_ALIGN_PARAGRAPH.CENTER)

for sec in doc.sections:
    f = sec.footer.paragraphs[0]
    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
    f.add_run('รายงาน PA ป.1 ฉบับหลักฐานจริง 2569 | หน้า ')
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    f._p.append(fld)

doc.core_properties.title = 'รายงาน PA ป1 ฉบับหลักฐานจริง 2569 นายอนันตชัย เพ็ชรรี่'
doc.core_properties.author = 'นายอนันตชัย เพ็ชรรี่'
doc.save(OUT)
print('SUCCESSFULLY_SAVED:', OUT)
